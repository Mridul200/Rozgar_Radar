import io
import re
from typing import Dict, Any, List, Tuple
import pypdf
import docx

EMAIL_REGEX = r"[a-zA-Z0-9_.+-]+@[a-zA-Z0-9-]+\.[a-zA-Z0-9-.]+"
PHONE_REGEX = r"(?:\+?\d{1,3}[-.\s]?)?\(?\d{3}\)?[-.\s]?\d{3}[-.\s]?\d{4}|\+91[-.\s]?[6-9]\d{9}"
LINKEDIN_REGEX = r"(?:https?:\/\/)?(?:www\.)?linkedin\.com\/in\/[a-zA-Z0-9_-]+"
GITHUB_REGEX = r"(?:https?:\/\/)?(?:www\.)?github\.com\/[a-zA-Z0-9_-]+"

SECTION_HEADERS = {
    "summary": ["professional summary", "summary", "career objective", "objective", "about me", "profile"],
    "experience": ["work experience", "experience", "employment history", "professional experience", "work history"],
    "education": ["education", "academic background", "qualifications", "academics"],
    "skills": ["technical skills", "skills", "core competencies", "skills & tools", "technologies"],
    "projects": ["projects", "personal projects", "academic projects", "key projects"],
    "certifications": ["certifications", "licenses", "certificates", "credentials"],
    "achievements": ["achievements", "honors", "awards", "accomplishments"]
}

def extract_text_from_pdf(file_bytes: bytes) -> Tuple[str, List[str]]:
    warnings = []
    text = ""
    try:
        reader = pypdf.PdfReader(io.BytesIO(file_bytes))
        num_pages = len(reader.pages)
        if num_pages == 0:
            warnings.append("The uploaded PDF has 0 pages.")
            return "", warnings

        for page_idx, page in enumerate(reader.pages):
            page_text = page.extract_text() or ""
            text += page_text + "\n"

        if len(text.strip()) < 50:
            warnings.append("Extracted text is very short. If this is a scanned image PDF, OCR may be needed.")
    except Exception as e:
        warnings.append(f"PDF extraction error: {str(e)}")

    return text.strip(), warnings

def extract_text_from_docx(file_bytes: bytes) -> Tuple[str, List[str]]:
    warnings = []
    text = ""
    try:
        doc = docx.Document(io.BytesIO(file_bytes))
        paras = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                paras.extend([c.text.strip() for c in row.cells if c.text.strip()])
        text = "\n".join(paras)
        if len(text.strip()) < 50:
            warnings.append("Extracted text is very short. Document might be empty or image-based.")
    except Exception as e:
        warnings.append(f"DOCX extraction error: {str(e)}")

    return text.strip(), warnings

def parse_resume_text_to_structure(raw_text: str) -> Dict[str, Any]:
    """Maps extracted raw text into Rozgar Radar standard ResumeData structure."""
    lines = [l.strip() for l in raw_text.splitlines() if l.strip()]
    if not lines:
        return {
            "fullName": "",
            "professionalTitle": "",
            "email": "",
            "phone": "",
            "location": "",
            "summary": "",
            "skills": {"technical": [], "tools": [], "soft": [], "other": []},
            "experience": [],
            "education": [],
            "projects": [],
            "certifications": [],
            "achievements": []
        }

    # 1. Contact information
    email_match = re.search(EMAIL_REGEX, raw_text)
    email = email_match.group(0) if email_match else ""

    phone_match = re.search(PHONE_REGEX, raw_text)
    phone = phone_match.group(0) if phone_match else ""

    linkedin_match = re.search(LINKEDIN_REGEX, raw_text, re.IGNORECASE)
    linkedin = linkedin_match.group(0) if linkedin_match else ""

    github_match = re.search(GITHUB_REGEX, raw_text, re.IGNORECASE)
    github = github_match.group(0) if github_match else ""

    # Name is typically on line 0 or 1 (excluding email/phone)
    full_name = ""
    for candidate in lines[:3]:
        if not re.search(EMAIL_REGEX, candidate) and not re.search(r"\d{5,}", candidate) and len(candidate.split()) in [1, 2, 3, 4]:
            full_name = candidate.strip()
            break
    if not full_name and lines:
        full_name = lines[0]

    # Target title or professional role
    prof_title = ""
    if len(lines) > 1 and lines[1] != full_name and len(lines[1]) < 60 and not re.search(EMAIL_REGEX, lines[1]):
        prof_title = lines[1]

    # Section segmenter
    sections: Dict[str, List[str]] = {}
    current_sec = "header"
    sections[current_sec] = []

    for line in lines:
        line_lower = line.lower().rstrip(":")
        matched_sec = None
        for sec, keywords in SECTION_HEADERS.items():
            if line_lower in keywords:
                matched_sec = sec
                break

        if matched_sec:
            current_sec = matched_sec
            if current_sec not in sections:
                sections[current_sec] = []
        else:
            sections[current_sec].append(line)

    # 2. Summary
    summary_lines = sections.get("summary", [])
    summary = " ".join(summary_lines)

    # 3. Skills
    skills_lines = sections.get("skills", [])
    skills_text = " ".join(skills_lines)
    # Split by commas, pipes, or bullets
    raw_skills = re.split(r"[,\|\•\-\/\;]", skills_text)
    technical_skills = []
    for s in raw_skills:
        clean = s.strip()
        if clean and len(clean) < 35 and clean.lower() not in ["and", "or", "skills", "tools"]:
            if clean not in technical_skills:
                technical_skills.append(clean)

    # 4. Experience
    exp_lines = sections.get("experience", [])
    experience = []
    if exp_lines:
        cur_exp = None
        for line in exp_lines:
            # Check if line looks like a title/company or date
            has_date = bool(re.search(r"\b(?:19|20)\d{2}\b|\bpresent\b|\bcurrent\b", line, re.IGNORECASE))
            if has_date or (cur_exp is None and len(line) < 80):
                if cur_exp:
                    experience.append(cur_exp)
                parts = re.split(r"[-–|@,]", line)
                title = parts[0].strip() if len(parts) > 0 else line
                comp = parts[1].strip() if len(parts) > 1 else ""
                cur_exp = {
                    "id": f"exp-{len(experience) + 1}",
                    "title": title,
                    "company": comp,
                    "location": "",
                    "startDate": "",
                    "endDate": "Present" if "present" in line.lower() else "",
                    "current": "present" in line.lower() or "current" in line.lower(),
                    "responsibilities": "",
                    "achievements": []
                }
            elif cur_exp:
                if line.startswith(("-", "•", "*")):
                    cur_exp["achievements"].append(line.lstrip("-•* ").strip())
                else:
                    if cur_exp["responsibilities"]:
                        cur_exp["responsibilities"] += " " + line
                    else:
                        cur_exp["responsibilities"] = line

        if cur_exp:
            experience.append(cur_exp)

    # 5. Education
    edu_lines = sections.get("education", [])
    education = []
    if edu_lines:
        cur_edu = None
        for line in edu_lines:
            has_year = bool(re.search(r"\b(?:19|20)\d{2}\b", line))
            if has_year or (cur_edu is None and len(line) < 90):
                if cur_edu:
                    education.append(cur_edu)
                year_match = re.search(r"\b(19\d{2}|20\d{2})\b", line)
                grad_year = year_match.group(0) if year_match else ""
                cur_edu = {
                    "id": f"edu-{len(education) + 1}",
                    "degree": line.strip(),
                    "school": "",
                    "location": "",
                    "gradYear": grad_year,
                    "grade": ""
                }
            elif cur_edu:
                if not cur_edu["school"]:
                    cur_edu["school"] = line.strip()
                elif "gpa" in line.lower() or "cgpa" in line.lower() or "%" in line:
                    cur_edu["grade"] = line.strip()

        if cur_edu:
            education.append(cur_edu)

    # 6. Projects
    proj_lines = sections.get("projects", [])
    projects = []
    if proj_lines:
        cur_proj = None
        for line in proj_lines:
            if cur_proj is None or (len(line) < 60 and not line.startswith(("-", "•"))):
                if cur_proj:
                    projects.append(cur_proj)
                cur_proj = {
                    "id": f"proj-{len(projects) + 1}",
                    "name": line.strip(),
                    "description": "",
                    "technologies": "",
                    "link": ""
                }
            elif cur_proj:
                if "http" in line:
                    url_m = re.search(r"https?:\/\/[^\s]+", line)
                    if url_m:
                        cur_proj["link"] = url_m.group(0)
                else:
                    if cur_proj["description"]:
                        cur_proj["description"] += " " + line
                    else:
                        cur_proj["description"] = line

        if cur_proj:
            projects.append(cur_proj)

    return {
        "fullName": full_name,
        "professionalTitle": prof_title,
        "email": email,
        "phone": phone,
        "location": "",
        "website": "",
        "linkedin": linkedin,
        "github": github,
        "summary": summary,
        "skills": {
            "technical": technical_skills[:20],
            "tools": [],
            "soft": [],
            "other": []
        },
        "experience": experience,
        "education": education,
        "projects": projects,
        "certifications": [],
        "achievements": sections.get("achievements", [])
    }


def parse_resume_file(file_bytes: bytes, filename: str) -> Dict[str, Any]:
    fn = filename.lower()
    raw_text = ''
    warnings = []
    
    if fn.endswith('.pdf'):
        raw_text, pdf_warnings = extract_text_from_pdf(file_bytes)
        warnings.extend(pdf_warnings)
    elif fn.endswith('.docx'):
        raw_text, docx_warnings = extract_text_from_docx(file_bytes)
        warnings.extend(docx_warnings)
    else:
        try:
            raw_text = file_bytes.decode('utf-8', errors='ignore')
        except Exception as e:
            warnings.append(f'Text decoding issue: {str(e)}')
            raw_text = ''

    parsed_data = parse_resume_text_to_structure(raw_text)
    
    if not parsed_data.get('email'):
        warnings.append('Could not extract email address automatically. Please verify Personal Info.')
    if not parsed_data.get('skills', {}).get('technical'):
        warnings.append('No technical skills were automatically extracted. You can add them in Skills step.')
    if not parsed_data.get('experience') and not parsed_data.get('education'):
        warnings.append('No experience or education blocks were clearly identified. Please review each section.')

    return {
        'raw_text': raw_text,
        'parsed_data': parsed_data,
        'warnings': warnings
    }
