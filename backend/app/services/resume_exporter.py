import io
from typing import Dict, Any, List
from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer,
    HRFlowable,
    ListFlowable,
    ListItem
)
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

def generate_pdf_resume(resume_data: Dict[str, Any], template: str = "classic_ats") -> bytes:
    """Generates an ATS-compliant, selectable-text PDF document from structured resume data."""
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=letter,
        rightMargin=0.55 * inch,
        leftMargin=0.55 * inch,
        topMargin=0.55 * inch,
        bottomMargin=0.55 * inch
    )

    styles = getSampleStyleSheet()

    # Define ATS-friendly typographic styles
    name_style = ParagraphStyle(
        "ResumeName",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=18,
        leading=22,
        alignment=0 if template == "technical" else 1, # Center or Left
        textColor=colors.HexColor("#0f172a" if template != "technical" else "#0284c7")
    )

    title_style = ParagraphStyle(
        "ResumeTitle",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=11,
        leading=14,
        alignment=0 if template == "technical" else 1,
        textColor=colors.HexColor("#475569")
    )

    contact_style = ParagraphStyle(
        "ResumeContact",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=9,
        leading=13,
        alignment=0 if template == "technical" else 1,
        textColor=colors.HexColor("#334155")
    )

    section_heading = ParagraphStyle(
        "SectionHeading",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=12,
        leading=15,
        spaceBefore=8,
        spaceAfter=3,
        textColor=colors.HexColor("#0f172a")
    )

    item_title = ParagraphStyle(
        "ItemTitle",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=10,
        leading=13,
        textColor=colors.HexColor("#1e293b")
    )

    item_subtitle = ParagraphStyle(
        "ItemSubtitle",
        parent=styles["Normal"],
        fontName="Helvetica-Oblique",
        fontSize=9.5,
        leading=12,
        textColor=colors.HexColor("#475569")
    )

    body_style = ParagraphStyle(
        "BodyTextCustom",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=9.5,
        leading=13,
        textColor=colors.HexColor("#1e293b")
    )

    bullet_style = ParagraphStyle(
        "BulletCustom",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=9,
        leading=12.5,
        leftIndent=12,
        firstLineIndent=-8,
        textColor=colors.HexColor("#1e293b")
    )

    story = []

    # 1. Header (Name & Contact)
    full_name = resume_data.get("fullName", "").strip() or "Candidate Resume"
    story.append(Paragraph(full_name, name_style))

    prof_title = resume_data.get("professionalTitle", "").strip()
    if prof_title:
        story.append(Paragraph(prof_title, title_style))

    contacts = []
    if resume_data.get("email"): contacts.append(resume_data["email"])
    if resume_data.get("phone"): contacts.append(resume_data["phone"])
    if resume_data.get("location"): contacts.append(resume_data["location"])
    if resume_data.get("linkedin"): contacts.append(resume_data["linkedin"])
    if resume_data.get("github"): contacts.append(resume_data["github"])
    if resume_data.get("website"): contacts.append(resume_data["website"])

    if contacts:
        contact_line = "  •  ".join(contacts)
        story.append(Spacer(1, 3))
        story.append(Paragraph(contact_line, contact_style))

    story.append(Spacer(1, 6))
    story.append(HRFlowable(width="100%", thickness=1, color=colors.HexColor("#cbd5e1"), spaceBefore=2, spaceAfter=6))

    # Helper to add section title
    def add_section(title: str):
        story.append(Paragraph(title.upper(), section_heading))
        story.append(HRFlowable(width="100%", thickness=0.5, color=colors.HexColor("#e2e8f0"), spaceBefore=1, spaceAfter=4))

    # 2. Professional Summary
    summary = resume_data.get("summary", "").strip()
    if summary:
        add_section("Professional Summary")
        story.append(Paragraph(summary, body_style))
        story.append(Spacer(1, 6))

    # 3. Technical Skills
    skills_data = resume_data.get("skills", {})
    all_skills = []
    if isinstance(skills_data, dict):
        tech = skills_data.get("technical", [])
        tools = skills_data.get("tools", [])
        soft = skills_data.get("soft", [])
        other = skills_data.get("other", [])
        if tech: all_skills.append(f"<b>Technical:</b> {', '.join(tech)}")
        if tools: all_skills.append(f"<b>Tools:</b> {', '.join(tools)}")
        if soft: all_skills.append(f"<b>Core Strengths:</b> {', '.join(soft)}")
        if other: all_skills.append(f"<b>Additional:</b> {', '.join(other)}")
    elif isinstance(skills_data, list) and skills_data:
        all_skills.append(f"<b>Skills:</b> {', '.join(skills_data)}")

    if all_skills:
        add_section("Skills & Technologies")
        for sk in all_skills:
            story.append(Paragraph(sk, body_style))
        story.append(Spacer(1, 6))

    # 4. Work Experience
    experience = resume_data.get("experience", [])
    if experience:
        add_section("Work Experience")
        for exp in experience:
            title = exp.get("title", "")
            company = exp.get("company", "")
            dates = f"{exp.get('startDate', '')} - {'Present' if exp.get('current') else exp.get('endDate', '')}".strip(" -")
            loc = exp.get("location", "")

            heading_parts = []
            if title and company:
                heading_parts.append(f"{title} | {company}")
            elif title or company:
                heading_parts.append(title or company)
            if loc:
                heading_parts.append(loc)
            if dates:
                heading_parts.append(dates)

            story.append(Paragraph(" — ".join(heading_parts), item_title))

            if exp.get("responsibilities"):
                story.append(Paragraph(exp["responsibilities"], body_style))

            achievements = exp.get("achievements", [])
            if isinstance(achievements, list):
                for bullet in achievements:
                    if bullet.strip():
                        story.append(Paragraph(f"• {bullet.strip()}", bullet_style))
            elif isinstance(achievements, str) and achievements.strip():
                story.append(Paragraph(f"• {achievements.strip()}", bullet_style))

            story.append(Spacer(1, 4))
        story.append(Spacer(1, 4))

    # 5. Projects
    projects = resume_data.get("projects", [])
    if projects:
        add_section("Key Projects")
        for proj in projects:
            p_name = proj.get("name", "")
            tech = proj.get("technologies", "")
            link = proj.get("link", "")

            p_title = p_name
            if tech:
                p_title += f" ({tech})"
            if link:
                p_title += f" [Link: {link}]"

            story.append(Paragraph(p_title, item_title))
            if proj.get("description"):
                story.append(Paragraph(proj["description"], body_style))
            if proj.get("outcomes"):
                story.append(Paragraph(f"• <b>Impact:</b> {proj['outcomes']}", bullet_style))
            story.append(Spacer(1, 4))
        story.append(Spacer(1, 4))

    # 6. Education
    education = resume_data.get("education", [])
    if education:
        add_section("Education")
        for edu in education:
            deg = edu.get("degree", "")
            school = edu.get("school", "")
            year = edu.get("gradYear", "")
            gpa = edu.get("grade") or edu.get("gpa") or ""

            edu_line = f"<b>{deg}</b>"
            if school: edu_line += f" — {school}"
            if year: edu_line += f" ({year})"
            if gpa: edu_line += f" | Grade/GPA: {gpa}"

            story.append(Paragraph(edu_line, body_style))
            if edu.get("coursework"):
                story.append(Paragraph(f"Coursework: {edu['coursework']}", item_subtitle))
            story.append(Spacer(1, 3))
        story.append(Spacer(1, 4))

    # 7. Certifications
    certs = resume_data.get("certifications", [])
    if certs:
        add_section("Certifications")
        for c in certs:
            c_name = c.get("name", "")
            issuer = c.get("issuer") or c.get("organization") or ""
            year = c.get("year") or c.get("issueDate") or ""
            cred_id = c.get("credentialId", "")

            line = f"<b>{c_name}</b>"
            if issuer: line += f" — {issuer}"
            if year: line += f" ({year})"
            if cred_id: line += f" [ID: {cred_id}]"

            story.append(Paragraph(f"• {line}", bullet_style))
        story.append(Spacer(1, 4))

    # 8. Achievements
    achievements = resume_data.get("achievements", [])
    if achievements:
        add_section("Honors & Achievements")
        for a in achievements:
            if isinstance(a, str) and a.strip():
                story.append(Paragraph(f"• {a.strip()}", bullet_style))
        story.append(Spacer(1, 4))

    doc.build(story)
    return buffer.getvalue()

def generate_docx_resume(resume_data: Dict[str, Any], template: str = "classic_ats") -> bytes:
    """Generates an ATS-compliant Word DOCX document from structured resume data."""
    doc = docx.Document()

    # Set page margins to 0.6 inch
    sections = doc.sections
    for s in sections:
        s.top_margin = Inches(0.6)
        s.bottom_margin = Inches(0.6)
        s.left_margin = Inches(0.6)
        s.right_margin = Inches(0.6)

    # 1. Header
    full_name = resume_data.get("fullName", "").strip() or "Candidate Resume"
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_after = Pt(2)
    run_name = title_p.add_run(full_name)
    run_name.bold = True
    run_name.font.size = Pt(18)
    run_name.font.name = "Arial"

    prof_title = resume_data.get("professionalTitle", "").strip()
    if prof_title:
        p_sub = doc.add_paragraph()
        p_sub.paragraph_format.space_after = Pt(4)
        run_title = p_sub.add_run(prof_title)
        run_title.font.size = Pt(11)
        run_title.font.color.rgb = RGBColor(71, 85, 105)

    contacts = []
    if resume_data.get("email"): contacts.append(resume_data["email"])
    if resume_data.get("phone"): contacts.append(resume_data["phone"])
    if resume_data.get("location"): contacts.append(resume_data["location"])
    if resume_data.get("linkedin"): contacts.append(resume_data["linkedin"])
    if resume_data.get("github"): contacts.append(resume_data["github"])

    if contacts:
        p_con = doc.add_paragraph()
        p_con.paragraph_format.space_after = Pt(8)
        run_con = p_con.add_run("  |  ".join(contacts))
        run_con.font.size = Pt(9.5)
        run_con.font.color.rgb = RGBColor(51, 65, 85)

    def add_heading(text: str):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(8)
        h.paragraph_format.space_after = Pt(3)
        run = h.add_run(text.upper())
        run.bold = True
        run.font.size = Pt(11)
        run.font.color.rgb = RGBColor(15, 23, 42)

    # 2. Professional Summary
    summary = resume_data.get("summary", "").strip()
    if summary:
        add_heading("Professional Summary")
        p = doc.add_paragraph(summary)
        p.paragraph_format.space_after = Pt(6)

    # 3. Skills
    skills_data = resume_data.get("skills", {})
    all_skills = []
    if isinstance(skills_data, dict):
        for k, v in skills_data.items():
            if v: all_skills.append(f"{k.title()}: {', '.join(v)}")
    elif isinstance(skills_data, list) and skills_data:
        all_skills.append(f"Skills: {', '.join(skills_data)}")

    if all_skills:
        add_heading("Skills & Technologies")
        for line in all_skills:
            p = doc.add_paragraph(line)
            p.paragraph_format.space_after = Pt(2)

    # 4. Work Experience
    experience = resume_data.get("experience", [])
    if experience:
        add_heading("Work Experience")
        for exp in experience:
            title = exp.get("title", "")
            company = exp.get("company", "")
            dates = f"{exp.get('startDate', '')} - {'Present' if exp.get('current') else exp.get('endDate', '')}".strip(" -")
            loc = exp.get("location", "")

            head_text = f"{title} at {company}" if (title and company) else (title or company)
            if loc: head_text += f" ({loc})"
            if dates: head_text += f" | {dates}"

            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run(head_text)
            run.bold = True

            if exp.get("responsibilities"):
                doc.add_paragraph(exp["responsibilities"])

            for bullet in exp.get("achievements", []):
                if isinstance(bullet, str) and bullet.strip():
                    doc.add_paragraph(bullet.strip(), style="List Bullet")

    # 5. Key Projects
    projects = resume_data.get("projects", [])
    if projects:
        add_heading("Key Projects")
        for proj in projects:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(1)
            r = p.add_run(proj.get("name", ""))
            r.bold = True
            if proj.get("technologies"):
                r2 = p.add_run(f" — Tech: {proj['technologies']}")
                r2.italic = True
            if proj.get("link"):
                p.add_run(f" [{proj['link']}]")

            if proj.get("description"):
                doc.add_paragraph(proj["description"])
            if proj.get("outcomes"):
                doc.add_paragraph(f"Impact: {proj['outcomes']}", style="List Bullet")

    # 6. Education
    education = resume_data.get("education", [])
    if education:
        add_heading("Education")
        for edu in education:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(1)
            deg = edu.get("degree", "")
            school = edu.get("school", "")
            year = edu.get("gradYear", "")
            r = p.add_run(f"{deg} — {school} ({year})")
            r.bold = True
            if edu.get("grade"):
                p.add_run(f" | Grade: {edu['grade']}")

    # 7. Certifications
    certs = resume_data.get("certifications", [])
    if certs:
        add_heading("Certifications")
        for c in certs:
            doc.add_paragraph(f"{c.get('name', '')} — {c.get('issuer', '')} ({c.get('year', '')})", style="List Bullet")

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()

